#!/usr/bin/env python3
"""Build the final COM7022 report (.docx and .pdf) from the master source _brain/REPORT.md.

Same content and visual language as the GitHub Pages portal: designed cover with KPI strip, contents page with
page numbers, part banners, Arden Midnight/Tiffany tables, callouts, recommendation cards, rendered equations,
and a caption under every figure/table/screenshot that links to its file in the GitHub repository.
Run:  python "COM7022 - Assignment/scripts/04_build_report_docx.py"   (needs LibreOffice for the PDF)
"""
import re
import subprocess
import tempfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

import report_source as rs

DOCX = rs.ROOT / "COM7022_Portfolio_25199053_final.docx"
PDF = DOCX.with_suffix(".pdf")
EQ_DIR = Path(tempfile.mkdtemp())

MIDNIGHT, MIDNIGHT_SOFT = RGBColor(0x00, 0x2B, 0x49), RGBColor(0x1A, 0x4A, 0x6B)
TIFFANY, TIFFANY_DEEP = RGBColor(0x66, 0xC9, 0xBA), RGBColor(0x3A, 0xA8, 0x9A)
TEAL, MUTE, INK, WHITE = RGBColor(0x0E, 0x7C, 0x8B), RGBColor(0x5A, 0x6B, 0x75), RGBColor(0x24, 0x35, 0x42), RGBColor(255, 255, 255)
HEX = {"midnight": "002B49", "soft": "1A4A6B", "tiffany": "66C9BA", "deep": "3AA89A", "teal": "0E7C8B",
       "mist": "E6F5F2", "paper": "F7FAFB", "grid": "D5DEE3", "red": "D9534F"}
FONT = "Calibri"


# ----------------------------------------------------------------------------- low-level helpers
def shade(cell, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(sh)


def cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (size in eighths of a point, hex colour); omitted edges get no border."""
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec[0])); el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    cell._tc.get_or_add_tcPr().append(b)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); m.append(el)
    cell._tc.get_or_add_tcPr().append(m)


def table_no_borders(table):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "nil"); b.append(el)
    table._tbl.tblPr.append(b)


def keep_rows_together(table):
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def para_border_bottom(p, color, size=12, space=4):
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
    bdr.append(el)
    p._p.get_or_add_pPr().append(bdr)


def spacer(doc, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 0.8
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def hyperlink(p, url, text, size=8, color=TEAL, bold=False):
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                            is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), str(color)); rPr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    if bold:
        rPr.append(OxmlElement("w:b"))
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r)
    p._p.append(h)


def run(p, text, size=None, bold=False, italic=False, color=None, font=None):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if color is not None:
        r.font.color.rgb = color
    if font:
        r.font.name = font
    return r


def inline(p, text, size=None, color=None):
    """Render **bold**, *italic*, `code` and bare URLs into runs."""
    for part in re.split(r"(\*\*.+?\*\*|(?<![*\w])\*(?!\s).+?(?<!\s)\*(?![*\w])|`[^`]+`|https?://\S+[^\s.,;:)])", text):
        if not part:
            continue
        if part.startswith("**"):
            run(p, part[2:-2], size, bold=True, color=color or MIDNIGHT)
        elif part.startswith("`"):
            run(p, part[1:-1], (size or 10.5) - 1, color=MIDNIGHT_SOFT, font="Consolas")
        elif part.startswith("http"):
            hyperlink(p, part, part, size=size or 10.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run(p, part[1:-1], size, italic=True, color=color)
        else:
            run(p, part, size, color=color)


# ----------------------------------------------------------------------------- document set-up
def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.1)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size, st.font.color.rgb = FONT, Pt(10.5), INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(6), 1.15
    for name, size, color, before, after in (("Heading 1", 18, MIDNIGHT, 4, 8), ("Heading 2", 12.5, MIDNIGHT_SOFT, 12, 4),
                                             ("Heading 3", 11, MIDNIGHT_SOFT, 8, 3)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold, h.font.color.rgb = FONT, Pt(size), True, color
        rf = h.element.rPr.rFonts
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), FONT)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(att), None)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(after)
        h.paragraph_format.keep_with_next = True
    lb = doc.styles["List Bullet"]
    lb.font.name, lb.font.size = FONT, Pt(10.5)
    lb.paragraph_format.space_after = Pt(4)


def header_footer(section):
    section.different_first_page_header_footer = True
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(hp, "COM7022 Machine Learning  ·  Product Segmentation Portfolio  ·  ", 8, color=MUTE)
    run(hp, f"Student ID {rs.STUDENT_ID}", 8, bold=True, color=MIDNIGHT)
    para_border_bottom(hp, HEX["tiffany"], 6, 4)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "Page ", 8.5, color=MUTE)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = fp.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = MUTE
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        r._r.append(el)


# ----------------------------------------------------------------------------- components
def band(doc, fill, rows, pad=(260, 260, 320, 320), left_bar=None):
    """Full-width coloured box. rows = [(text, size, bold, colour, space_after), ...]."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_no_borders(t)
    c = t.cell(0, 0); c.width = Cm(16.8)
    shade(c, fill); cell_margins(c, *pad)
    cell_borders(c, left=(40, left_bar)) if left_bar else None
    for i, (text, size, bold, color, after) in enumerate(rows):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.05
        run(p, text, size, bold=bold, color=color)
    return t


def kpi_strip(doc):
    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_no_borders(t)
    for i, (title, val, desc, acc) in enumerate(rs.KPIS):
        c = t.cell(0, i); c.width = Cm(4.2)
        cell_borders(c, top=(30, HEX["deep"] if acc else HEX["midnight"]), left=(4, HEX["grid"]),
                     right=(4, HEX["grid"]), bottom=(4, HEX["grid"]))
        cell_margins(c, 140, 140, 140, 110)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        run(p, title.upper(), 7.5, bold=True, color=MUTE)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1)
        run(p, val, 20, bold=True, color=TIFFANY_DEEP if acc else MIDNIGHT)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        run(p, desc, 8, color=MUTE)
    spacer(doc, 4)


def source_line(p, rel):
    run(p, "Source: ", 8, bold=True, color=TEAL)
    hyperlink(p, rs.repo_url(rel), rs.REPO_DIR + rel, 8)


def caption(doc, label, text, rel=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(12)
    run(p, f"{label}. ", 9, bold=True, color=MIDNIGHT)
    inline(p, text, 9, MUTE)
    if rel:
        p.add_run("\n")
        source_line(p, rel)


def figure(doc, path, label, text, rel, width=15.6, max_height=21.5, code=False):
    w, h = Image.open(path).size
    if code:                                   # same type size for every code screenshot
        width = min(width, w / 78)
    width = min(width, max_height * w / h)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    caption(doc, label, text, rel)


def data_table(doc, df, label, cap, rel, hl=None, font=8.5):
    pc = doc.add_paragraph(); pc.paragraph_format.keep_with_next = True
    pc.paragraph_format.space_before, pc.paragraph_format.space_after = Pt(6), Pt(3)
    run(pc, f"{label}. ", 9, bold=True, color=MIDNIGHT); inline(pc, cap, 9, MIDNIGHT)
    rows, cols = df.shape
    t = doc.add_table(rows=rows + 1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_no_borders(t)
    for j, col in enumerate(df.columns):
        c = t.cell(0, j); shade(c, HEX["midnight"]); cell_margins(c, 70, 70, 100, 80)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        run(p, str(col), font, bold=True, color=WHITE)
    for i in range(rows):
        for j in range(cols):
            c = t.cell(i + 1, j); cell_margins(c, 55, 55, 100, 80)
            cell_borders(c, bottom=(4, HEX["grid"]), left=(24, HEX["deep"]) if (i == hl and j == 0) else None)
            if i == hl:
                shade(c, HEX["mist"])
            elif i % 2:
                shade(c, HEX["paper"])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            strong = j == 0 or i == hl
            run(p, rs.fmt(df.iat[i, j]), font, bold=strong, color=MIDNIGHT if strong else INK)
    keep_rows_together(t)
    for row in t.rows[:-1]:
        for c in row.cells:
            c.paragraphs[0].paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(10)
    if rel:
        source_line(p, rel)


def callout(doc, title, text):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_no_borders(t)
    c = t.cell(0, 0); shade(c, HEX["mist"]); cell_borders(c, left=(40, HEX["deep"])); cell_margins(c, 170, 170, 240, 200)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    run(p, f"{title}: ", 10.5, bold=True, color=MIDNIGHT); inline(p, text, 10.5)
    spacer(doc, 4)


def rec_cards(doc):
    colours = {"c2": HEX["deep"], "c0": HEX["soft"], "c1": HEX["teal"], "c3": HEX["red"]}
    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_no_borders(t)
    for k, (cls, tag, title, finding, action) in enumerate(rs.RECS):
        c = t.cell(k // 2, k % 2); c.width = Cm(8.4)
        shade(c, "FFFFFF" if k in (0, 3) else HEX["paper"])
        cell_borders(c, left=(40, colours[cls]), top=(4, HEX["grid"]), right=(4, HEX["grid"]), bottom=(4, HEX["grid"]))
        cell_margins(c, 160, 160, 220, 160)
        col = RGBColor.from_string(colours[cls])
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1); run(p, tag.upper(), 7.5, bold=True, color=col)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(4); run(p, title, 11.5, bold=True, color=MIDNIGHT)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        run(p, "Evidence: ", 9.5, bold=True, color=MIDNIGHT); inline(p, finding, 9.5)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        run(p, "Action: ", 9.5, bold=True, color=MIDNIGHT); inline(p, action, 9.5)
    keep_rows_together(t)
    spacer(doc, 6)


def equation(doc, latex, n):
    tex = latex.replace(r"\lVert", r"\Vert").replace(r"\rVert", r"\Vert")
    path = EQ_DIR / f"eq{n}.png"
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${tex}$", fontsize=16, color="#002B49")
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    w, h = Image.open(path).size
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; table_no_borders(t)
    c = t.cell(0, 0); shade(c, HEX["paper"]); cell_margins(c, 140, 140, 200, 200)
    cell_borders(c, top=(4, HEX["grid"]), bottom=(4, HEX["grid"]), left=(4, HEX["grid"]), right=(4, HEX["grid"]))
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), height=Cm(h / 300 * 2.54 * 0.8))   # rendered at 300 dpi, shown at 80%
    spacer(doc, 4)


def part_banner(doc, tag, title, words):
    page_break(doc)
    band(doc, HEX["midnight"], [(tag.upper(), 9, True, TIFFANY, 2), (title, 20, True, WHITE, 3),
                                (f"{words:,} words of body text", 9, False, RGBColor(0xD5, 0xDE, 0xE3), 0)],
         pad=(300, 300, 340, 320), left_bar=HEX["tiffany"])
    spacer(doc, 6)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    para_border_bottom(p, HEX["tiffany"], 12, 4)


# ----------------------------------------------------------------------------- cover and contents
def cover(doc, wc):
    spacer(doc, 24)
    band(doc, HEX["midnight"], [
        ("MACHINE LEARNING PORTFOLIO  ·  MSC DATA SCIENCE  ·  ARDEN UNIVERSITY BERLIN", 8.5, True, TIFFANY, 12),
        (rs.TITLE, 25, True, WHITE, 12),
        ("An end-to-end unsupervised machine learning investigation comparing K-Means, Ward hierarchical clustering and "
         "DBSCAN across 1,000 marketplace listings, translated into evidence-based actions for catalogue strategy.",
         10.5, False, RGBColor(0xD5, 0xDE, 0xE3), 0)], pad=(620, 620, 480, 440), left_bar=HEX["tiffany"])
    spacer(doc, 10)
    meta = [("Student ID", rs.STUDENT_ID), ("Module", "COM7022 Machine Learning · Assignment [5444]"),
            ("Assessment", "Portfolio · Task 1 Practical Project · Task 2 Critical Report"),
            ("Word count", f"{wc['total']:,} words of body text (Task 1: {wc['task1_total']:,} · Task 2: {wc['task2']:,}); "
                           "captions, tables, figures, code screenshots, references and appendix excluded"),
            ("Dataset", "Sales-[5444].csv · 1,000 listings × 10 attributes"),
            ("Repository", rs.REPO), ("Review portal", rs.PAGES)]
    t = doc.add_table(rows=len(meta), cols=2); table_no_borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        a, b = t.cell(i, 0), t.cell(i, 1)
        a.width, b.width = Cm(3.4), Cm(13.4)
        for c in (a, b):
            cell_borders(c, bottom=(4, HEX["grid"])); cell_margins(c, 80, 80, 60, 60)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(a.paragraphs[0], k.upper(), 8, bold=True, color=MUTE)
        if v.startswith("http"):
            hyperlink(b.paragraphs[0], v, v, 10, TEAL, True)
        else:
            run(b.paragraphs[0], v, 10, bold=(i == 0), color=MIDNIGHT)
    spacer(doc, 14)
    kpi_strip(doc)
    p = doc.add_paragraph()
    run(p, "Every figure, table and code screenshot carries a caption that names and links its exact file in the GitHub repository.",
        8.5, italic=True, color=MUTE)


def contents(doc, blocks, pages):
    page_break(doc)
    h1(doc, "Contents")
    t = doc.add_table(rows=0, cols=2); table_no_borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for b in blocks:
        if b["type"] == "part":
            cells = t.add_row().cells
            p = cells[0].paragraphs[0]; p.paragraph_format.space_before = Pt(10)
            run(p, f"{b['tag']}  ·  {b['title']}".upper(), 8.5, bold=True, color=TIFFANY_DEEP)
        elif b["type"] in ("h1", "h2"):
            lvl1 = b["type"] == "h1"
            cells = t.add_row().cells
            p = cells[0].paragraphs[0]
            p.paragraph_format.left_indent = Cm(0 if lvl1 else 0.7)
            p.paragraph_format.space_before = Pt(5 if lvl1 else 0)
            run(p, b["text"], 10.5 if lvl1 else 9.5, bold=lvl1, color=MIDNIGHT if lvl1 else INK)
            q = cells[1].paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            q.paragraph_format.space_before = Pt(5 if lvl1 else 0)
            run(q, str(pages.get(b["text"], "")), 10 if lvl1 else 9, bold=lvl1, color=MIDNIGHT)
            if lvl1:
                for c in cells:
                    cell_borders(c, bottom=(4, HEX["grid"]))
        else:
            continue
        for c, w in zip(cells, (15.2, 1.6)):
            c.width = Cm(w); cell_margins(c, 20, 20, 40, 40)
            c.paragraphs[0].paragraph_format.space_after = Pt(1)


# ----------------------------------------------------------------------------- build
def build(pages):
    blocks, labels, tables, wc = rs.load()
    doc = Document()
    setup(doc)
    header_footer(doc.sections[0])
    cover(doc, wc)
    contents(doc, blocks, pages)
    after_part, eq_n, section = False, 0, ""
    for b in blocks:
        t = b["type"]
        if t == "part":
            part_banner(doc, b["tag"], b["title"], wc["task1_total"] if "1" in b["tag"] else wc["task2"])
            after_part = True
        elif t == "h1":
            # sections flow on; only the back matter starts on a fresh page (parts have their own break)
            if b["text"].startswith(("References", "Appendix")):
                page_break(doc)
            elif not after_part:
                spacer(doc, 4)
            after_part = False
            section = b["text"]
            h1(doc, b["text"])
        elif t in ("h2", "h3"):
            doc.add_heading(b["text"], level=int(t[1]))
        elif t == "para":
            inline(doc.add_paragraph(), b["text"])
        elif t == "bullets" and section == "References":
            for x in b["items"]:                  # Harvard list: hanging indent, no bullets
                p = doc.add_paragraph(); inline(p, x, 9.5)
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.8), Cm(-0.8)
        elif t == "bullets":
            for x in b["items"]:
                inline(doc.add_paragraph(style="List Bullet"), x)
        elif t in ("figure", "shot"):
            path, rel = rs.figure_path(b)
            text = ("Code screenshot: " + b["caption"]) if t == "shot" else b["caption"]
            figure(doc, path, b["label"], text, rel, 15.6, code=(t == "shot"))
        elif t == "table":
            df, cap, rel = tables[b["key"]]
            data_table(doc, df, b["label"], cap, rel, rs.HIGHLIGHT.get(b["key"]))
        elif t == "callout":
            callout(doc, b["title"], b["text"])
        elif t == "equation":
            eq_n += 1
            equation(doc, b["latex"], eq_n)
        elif t == "recs":
            rec_cards(doc)
        elif t == "kpis":
            kpi_strip(doc)
        elif t == "appendix":
            current = None
            for head, n, path, rel in rs.appendix_cells():
                if head != current:
                    current = head
                    doc.add_heading(f"A · {head}", level=3)
                figure(doc, path, f"Code A{n}", f"Notebook code cell {n}", rel, 15.6, code=True)
        # 'explorer' is interactive and portal-only; the same numbers are in the cluster profile table
    doc.save(DOCX)
    return blocks, wc


def to_pdf():
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(DOCX.parent), str(DOCX)],
                   check=True, capture_output=True)


def heading_pages(blocks):
    """PDF page of every h1/h2 heading, searched in order after the contents page."""
    text = subprocess.run(["pdftotext", "-layout", str(PDF), "-"], capture_output=True, text=True, check=True).stdout
    pages = text.split("\f")
    norm = lambda s: re.sub(r"\s+", " ", s).strip().lower()
    found, start = {}, 2
    for b in blocks:
        if b["type"] in ("h1", "h2"):
            key = norm(b["text"])[:38]
            for i in range(start, len(pages)):
                if key in norm(pages[i]):
                    found[b["text"]] = i + 1
                    start = i
                    break
    return found


def main():
    blocks, wc = build({})
    to_pdf()
    pages = heading_pages(blocks)          # second pass writes real page numbers into the contents page
    build(pages)
    to_pdf()
    final = heading_pages(blocks)
    if final != pages:
        build(final); to_pdf()
    print("word count:", wc)
    print("saved", DOCX.name, "and", PDF.name)


if __name__ == "__main__":
    main()
